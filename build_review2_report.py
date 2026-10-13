"""Build the Review 2 report from checked project evidence."""
import csv
import json
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
DATA = ROOT / "results_paired_both_review"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "OS_Review2_Methodology_and_Partial_Implementation.docx"


def csv_rows(path):
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


summary = {(r["workload"], r["method"]): r for r in csv_rows(DATA / "summary.csv")}
manifest = json.loads((DATA / "manifest.json").read_text(encoding="utf-8"))
train_manifest = json.loads((ROOT / "results_final_linux" / "manifest.json").read_text(encoding="utf-8"))
assert manifest["status"] == "PASS" and manifest["raw_rows"] == 105
assert train_manifest["training"]["training_rows"] == 90

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.25)
section.right_margin = Cm(2.25)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.13
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
for name, size, before, after in (("Title", 17, 0, 5), ("Heading 1", 13.5, 12, 5), ("Heading 2", 11.5, 9, 4)):
    style = styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
styles["Title"].paragraph_format.line_spacing = 1.0
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_style_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_style_ppr.remove(title_border)


def set_font(run, size=None, bold=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    return run


def para(text="", *, bold_lead=None, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def heading(text, level=1):
    doc.add_heading(text, level=level)


def cell_border(cell, color="BFBFBF"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        child = borders.find(qn(tag))
        if child is None:
            child = OxmlElement(tag)
            borders.append(child)
        child.set(qn("w:val"), "single")
        child.set(qn("w:sz"), "4")
        child.set(qn("w:color"), color)


def cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tcPr.append(shade)


def cell_margin(cell, top=70, start=85, bottom=70, end=85):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcPr.append(mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = mar.find(qn("w:" + key))
        if el is None:
            el = OxmlElement("w:" + key)
            mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def table(headers, data, widths, size=9.3):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Cm(width)
    for i, header in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i])
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(str(header)), size=size, bold=True)
        cell_shading(c, "E8E8E8")
        cell_border(c)
        cell_margin(c)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    rep.set(qn("w:val"), "true")
    trPr.append(rep)
    for row in data:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            c.width = Cm(widths[i])
            c.text = ""
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(str(value)), size=size)
            cell_border(c)
            cell_margin(c)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def caption(text, keep_next=True):
    p = para(text)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = keep_next
    for run in p.runs:
        set_font(run, size=9)


def flow(label, steps):
    p = para(label, keep=True)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    for r in p.runs:
        set_font(r, size=10.3, bold=True)
    box = doc.add_table(rows=1, cols=len(steps) * 2 - 1)
    box.autofit = False
    widths = [Cm(2.55) if i % 2 == 0 else Cm(0.7) for i in range(len(steps) * 2 - 1)]
    for i, w in enumerate(widths):
        box.columns[i].width = w
        cell = box.cell(0, i)
        cell.width = w
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pp = cell.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.space_after = Pt(0)
        set_font(pp.add_run(steps[i // 2] if i % 2 == 0 else ">"), size=9, bold=i % 2 == 1)
        if i % 2 == 0:
            cell_border(cell)
            cell_margin(cell, 60, 45, 60, 45)
        else:
            cell_margin(cell, 0, 0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def screenshot(filename, crop, width_cm, text):
    original = DATA / "terminal_screenshots" / filename
    assert original.exists(), original
    im = Image.open(original)
    trimmed = im.crop(crop)
    trimmed_path = OUT / ("crop_" + filename)
    trimmed.save(trimmed_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(trimmed_path), width=Cm(width_cm))
    caption(text, keep_next=False)


def figure(filename, width_cm, text):
    caption(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(OUT / "figures" / filename), width=Cm(width_cm))


title = doc.add_paragraph(style="Title")
set_font(title.add_run("Machine-Learning-Assisted Adaptive Chunk Scheduling for Multiprocessing Workloads"), size=17, bold=True)
sub = para("Operating Systems Review 2  Methodology and Partial Implementation")
sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
for r in sub.runs:
    set_font(r, size=11.5, bold=True)
author = para("Sashwath J   24BAI0059")
author.alignment = WD_ALIGN_PARAGRAPH.LEFT
for r in author.runs:
    set_font(r, size=10.5)
para("This prototype sends independent CPU tasks to a Linux process pool. Fixed chunks of four tasks are the reference. Two alternatives are implemented: a decision tree changes chunk size using live measurements, while a cost-balanced dispatcher groups tasks using known metadata. The results show different strengths and costs. The combined method proposed for the next review is not yet implemented.")

heading("1 Problem and proposed method")
para("A fixed chunk size gives each worker the same number of tasks per batch even when task cost varies. Small chunks can create dispatch overhead; large chunks can leave one worker with a long final batch. The project tests two responses to this problem. The decision-tree path observes early execution and predicts one of four chunk sizes. The cost-balanced path uses available task metadata to make batches of roughly similar estimated cost, then sends longer batches first. The Linux kernel still schedules the worker processes; this project changes user-space task dispatch, not kernel code.")
para("The next question is whether a combined scheduler can use the tree's workload signal and cost balancing without the current tree path's repeated round barrier. This is a testable design, not a measured result. Machine-learning-assisted task scheduling and adaptive self-scheduling already exist in the literature (Qawasmeh et al., 2016; Booth & Lane, 2022; Nookala et al., 2024).")

heading("2 System architecture")
figure("architecture.png", 16.5, "Figure 1  Implemented user space scheduling architecture. The dashed line shows runtime feedback used only by the tree path.")
para("The controller runs in Python 3.12.3 on Ubuntu under WSL2. It creates tasks, chooses fixed, tree or cost-balanced dispatch, and sends batches to two worker processes. Each worker calls the compiled C11 CPU kernel through ctypes. The parent records timing and checksums after results return. Model fitting is offline; the measured run loads an already trained tree.")

heading("3 Major modules and workflows")
table(["Module", "Implemented responsibility", "Main output"], [
    ("Task generator", "Recreates 192 CPU tasks from a workload name and seed.", "Task list and seed"),
    ("Native worker", "Executes synthetic iteration work or prime counting in C.", "Checksum and service time"),
    ("Tree trainer", "Profiles candidate chunks, labels fastest choices, fits a depth-3 tree.", "model.json"),
    ("Tree scheduler", "Runs a pilot, measures feedback, predicts the next chunk per round.", "Decision trace and timing"),
    ("Cost balancer", "Estimates task cost, packs batches, sorts them longest first.", "Batches and timing"),
    ("Evaluator", "Times matched runs and audits checksums and summary medians.", "raw.csv and summary.csv"),
], [3.2, 9.35, 3.95], 9.2)
figure("scheduling_flow.png", 16.5, "Figure 2  The two implemented scheduling paths. The tree waits at feedback rounds; cost balancing prepares batches once.")
para("The generator first creates the same seeded task list for each method. Offline training profiles candidate chunks and saves model.json; the evaluator later times each method on matched inputs, checks output equality, and recomputes the paired statistics. These training and evaluation steps run outside the two dispatch paths shown in Figure 2.")

heading("4 Dataset and collection procedure")
para("The inputs are reproducible generated CPU workloads, not an external dataset or collected production trace. Every timing group uses the same task list for all three methods. Five general patterns each have five seeds (7001-7005); the late-heavy pattern has ten seeds (11001-11010). Each group contains 192 tasks on two workers. Training uses different seeds (11-16), the holdout uses 101-104, and the paired benchmark uses the later seeds. This separation limits direct train/test leakage.")
table(["Pattern", "How the 192 tasks are generated", "Why included"], [
    ("Uniform", "Each synthetic task has 1.5 million C iterations.", "Low variability control"),
    ("Moderate", "Each task scales 1.5 million iterations by a seeded factor from 0.5 to 1.5.", "Ordinary uneven work"),
    ("High", "70 percent are about 0.2 times and the remainder about 2.87 times base work.", "Strong size variation"),
    ("Phase shift", "First half is light; second half mixes very light and heavy work.", "Distribution change during a run"),
    ("Prime scan", "C counts primes in successive 3,000-integer ranges.", "Non-synthetic computational kernel"),
    ("Late heavy", "One of the last eight tasks has 40 times base iterations; others vary 0.9-1.1 times.", "Rare delayed straggler"),
], [2.35, 9.75, 4.4], 9.0)
para("The training table has 90 fitted rows and 60 held-out rows. Each row is a measurement at a live-equivalent decision boundary. Its five predictors are mean task service time (average worker C time per task), batch spread ((maximum minus minimum batch service time) divided by maximum), throughput (tasks completed divided by round seconds), current chunk size, and fraction of tasks remaining. The label is the fastest median candidate among chunks 1, 4, 16 and 64; choices within 3 percent of fastest prefer the larger chunk. These sizes span one task through coarse batching, rather than being arbitrary output classes. The held-out exact label accuracy is 73.3 percent (44 of 60); this is a classification result, not a runtime speedup.")
para("Collection is deterministic from the generator and saved seeds. The earlier full run saved 4,800 task records in results_final_linux/dataset.csv. The paired study stores each group seed, method order, full timings and checksum in raw.csv; its generator code recreates the exact task list. No personally identifiable or third-party data is used.")

heading("5 Tools technologies and algorithms")
para("Execution used Ubuntu on WSL2, Python 3.12.3 for the controller and multiprocessing pool, GCC 13.3 for the C11 work kernel, ctypes for Python-to-C calls, and the Python standard library for a custom CART-style decision tree and CSV/JSON output. The fixed reference groups four tasks. The tree is depth 3 with minimum leaf size 5. The cost method sets target cost to total estimated cost divided by 32, greedily packs adjacent tasks up to 32 items, sorts batches by estimated cost and dispatches them once. Synthetic cost is estimated from the known iteration count. Prime cost uses work multiplied by an approximate square-root scale of the number range. The cost estimate is task metadata, not a learned prediction.")

heading("6 Experimental plan and measures")
para("The planned comparison is fixed chunk 4 versus online decision tree versus cost-balanced 32 under matched tasks. For each workload and seed, method order is seeded and shuffled. The primary measure is the median across seeds of each method's wall-time ratio to fixed 4 on the same seed. A ratio below 1 means lower full elapsed time; 0.870 means about 13 percent less time. Wall time starts before batch planning or pool creation and stops after results return and pool teardown. Offline model training is excluded from a test run but reported separately. This boundary measures application-level execution, not a kernel scheduling event.")
table(["Measure", "Definition", "Interpretation"], [
    ("Full wall time", "Time from scheduling start through planning, pool startup, dispatch and teardown.", "Primary user-visible latency; lower is better."),
    ("Paired ratio", "Method time divided by fixed-4 time for the same input seed; median of ratios.", "Below 1 faster; above 1 slower than baseline."),
    ("Batch p95", "Per-run 95th percentile from batch submission to worker completion.", "Shows slow-tail behavior; not the same as total wall time."),
    ("Startup", "Time from run start until worker pool is ready.", "Part of full wall time."),
    ("Checksum", "Sum of deterministic C results for all 192 tasks.", "Equality checks output correctness across methods."),
], [2.4, 8.45, 5.65], 9.0)
para("The benchmark has 35 matched input groups and 105 timed method runs. Its saved audit recomputes all summary medians, checks method coverage and order, and confirms equal checksums in each group. The sample is exploratory on one WSL2 machine with two logical CPUs. No confidence interval or cross-machine generalization is claimed.")

comparison_heading = doc.add_heading("7 Measured comparison and interpretation", level=1)
comparison_heading.paragraph_format.page_break_before = True
caption("Table 1  Full wall time and median paired ratio by workload")
results = []
for kind in ("uniform", "moderate", "high", "phase_shift", "prime_scan", "late_heavy"):
    f, m, b = (summary[(kind, method)] for method in ("fixed_4", "ml", "balanced_32"))
    results.append((kind.replace("_", " "), f["runs"], f"{float(f['median_wall_s']):.3f}",
                    f"{float(m['median_wall_s']):.3f} / {float(m['median_paired_ratio']):.3f}",
                    f"{float(b['median_wall_s']):.3f} / {float(b['median_paired_ratio']):.3f}"))
table(["Workload", "n", "Fixed 4 s", "Tree s / ratio", "Balanced s / ratio"], results,
      [3.4, 0.8, 2.45, 4.65, 5.2], 9.1)
para("Values are medians over five seeds per general workload and ten late-heavy seeds. The ratio is the median of within-seed ratios, so it need not equal a ratio of the two displayed median seconds. This matters for moderate and prime scan, where comparing the seconds alone could give a different ranking from the paired statistic.")
figure("paired_ratio_plot.png", 16.5, "Figure 3  Paired wall-time ratios calculated directly from the saved benchmark summary. The vertical line is the fixed-4 baseline.")
para("The tree improves the paired ratio on prime scan (0.894) but is slower than fixed 4 on uniform, moderate, high, phase-shift and late-heavy tasks. Its high-variance ratio is 1.245, about 24.5 percent more time. A plausible contributor is the repeated round barrier in the current implementation: each model decision waits for a round's submitted batches to finish. Prediction itself takes only microseconds in the live trace, but the full dispatch pattern can still be slower. The present data do not isolate a single cause.")
para("Cost balancing is strongest on the intended late-heavy pattern: median paired ratio 0.870, or about 13 percent less full time than fixed 4; median wall times are 0.432 s versus 0.496 s. It also improves the paired ratio on uniform, moderate, phase-shift and prime scan, but is slightly slower on high variance (1.006). This is evidence for a targeted benefit, not universal superiority. It does not make the cost method machine learning.")
doc.add_page_break()
caption("Table 2  Late heavy diagnostics  Ten matched seeds")
late = []
for label, method in (("Fixed 4", "fixed_4"), ("Decision tree", "ml"), ("Cost balanced 32", "balanced_32")):
    row = summary[("late_heavy", method)]
    late.append((label, f"{float(row['median_wall_s']):.3f}",
                 f"{float(row['median_paired_ratio']):.3f}",
                 f"{float(row['median_batch_p95_s']):.3f}",
                 f"{float(row['median_startup_s']):.3f}"))
table(["Method", "Wall s", "Paired ratio", "Batch p95 s", "Startup s"], late,
      [4.5, 2.6, 2.75, 3.5, 3.15], 9.2)
figure("late_heavy_plot.png", 14.8, "Figure 4  Late-heavy full wall time and batch p95 from ten matched seeds. The two measures describe different timing boundaries.")
para("A lower batch p95 does not guarantee lower total wall time. In the late-heavy group the tree has the lowest batch p95 (0.156 s) but slower total time (0.522 s). The cost path starts its estimated heavy batch early; in the live demonstration, task 187 moved to dispatch position 0. Its batch p95 is higher (0.412 s) while total time is lower. The single live run itself is not used as the ten-seed effect estimate; its exact elapsed times can change with system load.")

heading("8 Implementation progress and valid outputs")
para("Six of seven planned modules work: task generation, compiled C execution, tree training, online tree scheduling, cost-balanced scheduling, and paired evaluation. The combined overhead-aware scheduler remains for the next stage. Module count is not a precise measure of effort, but the working prototype exceeds the Review 2 half-complete requirement. No hybrid benchmark is claimed.")
table(["Delivered module", "Evidence to show", "Status"], [
    ("Generator and C kernel", "workload.py; native/work_kernel.c; repeatable checksums", "Complete"),
    ("Tree training", "training_data.csv; model.json; 73.3 percent held-out accuracy", "Complete"),
    ("Online tree run", "live demo trace: chunk 4 to 16; equal checksums", "Complete"),
    ("Cost-balanced run", "live demo: late task 187 dispatched first; equal checksums", "Complete"),
    ("Paired evaluator", "105 raw rows; 18 summary rows; all groups audited", "Complete"),
    ("Automated checks", "12 unit tests passing in Ubuntu", "Complete"),
    ("Combined scheduler", "Design and ablation pending", "Remaining"),
], [4.1, 9.9, 2.5], 9.0)
heading("9 Challenges and remaining work")
para("Chunk-label accuracy did not predict wall time; the tree's round barrier adds waiting, while cost balancing needs useful task metadata. Synthetic tasks and one two-worker WSL2 host limit generalization.")
para("Next, test an overhead-gated combined scheduler: use the tree after a pilot, balance only unstarted tasks when metadata are informative, and remove unnecessary round barriers. Compare fixed 4, tree, cost balance and the combined path on matched seeds. Then perturb cost estimates, test more workers and longer tasks, report uncertainty intervals, and repeat on another Linux machine or real trace. Keep the combined path only if full wall time improves with equal outputs.")

heading("References")
refs = [
    "Booth, J. D., & Lane, P. A. (2022). An adaptive self-scheduling loop scheduler. Concurrency and Computation: Practice and Experience, 34(6), e6750. https://doi.org/10.1002/cpe.6750",
    "Nookala, P., Chard, K., & Raicu, I. (2024). X-OpenMP: eXtreme fine-grained tasking using lock-less work stealing. Future Generation Computer Systems, 159, 444-458. https://doi.org/10.1016/j.future.2024.05.019",
    "Qawasmeh, A., Malik, A. M., & Chapman, B. M. (2016). Adaptive OpenMP task scheduling using runtime APIs and machine learning. In Proceedings of the 2015 IEEE 14th International Conference on Machine Learning and Applications (pp. 889-895). https://doi.org/10.1109/ICMLA.2015.111",
]
for ref in refs:
    p = para(ref)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    p.paragraph_format.space_after = Pt(6)

heading("Appendix  Terminal evidence")
para("The WSL project has a logical directory alias at the prompt shown below. The screenshots are cropped for print readability; the white-on-black text is the original terminal display.")
screenshot("dark_01_tests.png", (12, 12, 1290, 170), 16.0,
           "Figure A1  Fresh Linux unit-test run: 12 tests, OK.")
screenshot("dark_02_audit.png", (12, 12, 1190, 365), 16.0,
           "Figure A2  Saved benchmark audit: 35 paired groups, 105 runs, equal checksums and recalculated ratios.")
screenshot("dark_05_both.png", (12, 12, 1260, 265), 16.0,
           "Figure A3  Fresh live runs of both paths: the tree changes chunk 4 to 16; cost balancing dispatches the late task first. Both pass checksum checks. The tree trace is filtered to its first decisions and final result.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Page "), size=9)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.save(DOCX)
print(DOCX)
