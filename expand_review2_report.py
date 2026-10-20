"""Expand the edited Review 2 report with verified project details."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "output" / "OS_Review2_User_Edited_With_Module_Workflows.docx"
TARGET = ROOT / "output" / "OS_Review2_Expanded_Report.docx"


INSERTIONS = [
    (
        "System architecture",
        [
            ("h", "1.1 What the scheduler controls"),
            ("p", "The unit of work is one independent CPU task. A task contains an amount of work, a deterministic seed or starting integer, and a kernel type. A chunk is a group of these tasks sent to one process-pool worker as one batch. Changing the chunk changes how often the parent submits batches and how much work can be left in a worker's final batch. It does not change the Linux CPU time-sharing policy. The operating system still decides when each worker process runs; the project decides which application tasks are grouped and submitted to those processes."),
            ("p", "With 192 tasks, fixed chunk 4 forms 48 four-task batches. A small chunk gives the pool more chances to distribute uneven work but creates more submissions and result handling. A large chunk reduces that administrative work, yet a costly task inside a late batch can determine when the entire job finishes. The correct choice therefore depends on both task variability and scheduling overhead. The benchmark tests that tradeoff at the complete application boundary instead of assuming that fewer batches must always be faster."),
            ("h", "1.2 Scope of the claim"),
            ("p", "The work has two implemented alternatives to the fixed-4 reference: a trained tree that selects a later chunk from runtime measurements, and a separate metadata-based cost balancer. Both execute the same generated tasks and must return the same checksum. The planned combined scheduler is a research direction, not a completed result. This distinction matters because the measured cost-balanced improvement cannot be presented as an ML improvement, and a model's ability to predict a profiling label cannot be treated as proof of lower wall time."),
        ],
    ),
    (
        "Major modules and workflows",
        [
            ("h", "2.1 Where time is spent"),
            ("p", "The parent process constructs the task list, selects the scheduling method, creates a pool of two worker processes, submits batches, receives results, and closes the pool. Workers load the compiled C library through ctypes when needed. In the synthetic kernel, each task performs a deterministic integer loop; in the prime-scan kernel, it counts primes in its assigned integer range. The returned value lets the parent compare scheduler outputs without depending only on timing. The C code makes the task computation substantially different from a Python loop, but Python still owns scheduling and data movement."),
            ("p", "The main wall-time timer begins before method-specific planning or pool creation and ends after results return and the pool closes. That includes the costs a caller experiences when invoking this program. Offline model fitting is performed separately and is not hidden inside each timed invocation. Pool startup, dispatch, round waiting, native computation, result collection, and teardown can all influence the full time. Consequently, a microsecond-scale tree decision can coexist with a slower overall tree run."),
        ],
    ),
    (
        "Dataset and collection procedure",
        [
            ("h", "3.1 The task and worker path"),
            ("p", "The generator in os_sched/workload.py creates frozen task records from a workload name and seed. For the five general patterns, the paired review uses 192 tasks per record and passes either 1.5 million base C iterations or a 3,000-integer prime range. The late-heavy generator in os_sched/straggler_workload.py places one 40-times-cost task among the last eight task positions; the remaining tasks fluctuate between 0.9 and 1.1 times base work. This is intentionally constructed to test a late straggler, not presented as an observed production workload."),
            ("p", "The parent converts the task list into batches. Each worker applies the chosen C function to every task in its batch, sums the deterministic results, and records the batch's start, end, and service time. The parent then sums batch checksums. If fixed, tree, and balanced paths disagree on that final checksum for the same seed, the comparison is rejected. The test suite also checks generator determinism and the scheduling modules. A matching checksum is an output-equivalence check for these runs, not a proof that every possible task ordering is correct."),
            ("h", "3.2 How the tree learns a chunk label"),
            ("p", "Offline profiling first runs a pilot at a live-equivalent decision boundary. For each of five workload types, six training seeds and three current chunks (1, 4, and 16) produce 90 fitted records. Four separate holdout seeds produce 60 more records. After the pilot, the same next 64 tasks are tested with candidate chunks 1, 4, 16, and 64. Each candidate is measured twice, and candidate order is shuffled in each pass to reduce systematic order effects. The label is the chunk with the lowest candidate median; if several candidates are within 3 percent of the fastest, the larger chunk is preferred to avoid needless small-batch work."),
            ("p", "The five inputs saved with each label are mean measured task service time, normalized spread of batch service times, completed tasks per second in the pilot, the current chunk, and the fraction of tasks still unstarted. The model is a depth-3 CART-style classifier with at least five records per leaf. Its JSON file stores threshold tests and leaf labels, so a live run can load the fitted tree without an external ML package. On the separate 60-row holdout, 44 labels match exactly, or 73.3 percent. That number describes agreement with profiling labels only; it says nothing by itself about the full application's speed."),
            ("h", "3.3 What happens during a live tree run"),
            ("p", "The tree path begins with chunk 4. In each round it selects up to current chunk multiplied by workers multiplied by two unstarted tasks; the first two-worker round therefore submits 16 tasks. The pool call completes that round before the parent computes throughput, batch spread, mean task service time, and the remaining fraction. The saved tree predicts the chunk for the next round. A decision event records the old and new chunks, reason, elapsed round time, and inference time. The live evidence shows a 4-to-16 prediction, but that trace is one demonstration rather than an aggregate performance estimate."),
            ("p", "Waiting for a round to finish provides coherent feedback, but it is also a synchronization barrier. A fast worker may be idle while a slower batch determines when the parent can decide again. That limitation is visible in the distinction between the tree's low batch p95 and its higher full wall time on the late-heavy case. The program also contains a guarded-tree variant that can reduce a predicted chunk after a measured task-cost rise; that earlier variant is not one of the three methods in the current paired Review 2 table."),
            ("h", "3.4 What cost balancing does instead"),
            ("p", "The balanced path estimates synthetic cost from iterations and prime cost from range length times an approximate square-root factor for integer magnitude. It divides total estimated work by a target of 32, groups adjacent tasks until the target-cost or 32-item limit is met, then sorts batches from highest to lowest estimated cost. Thirty-two is a target, not a promise of exactly 32 batches."),
            ("p", "The pool receives the prepared list once, starting longer estimated batches early. The live late-heavy run places task 187 in dispatch position 0. This is not a learned model and cannot react to work whose cost is unknown or badly estimated. It complements the online tree, but a hybrid is still unmeasured."),
        ],
    ),
    (
        "Tools technologies and algorithms",
        [
            ("h", "4.1 What is saved and what can be reconstructed"),
            ("p", "The project does not import a public scheduling dataset. Inputs are generated deterministically from workload definitions and seeds. The earlier full training run saves task-level rows in results_final_linux/dataset.csv and the fitted tree in model.json. The later paired review uses different seeds and saves one timing row per method and seed in results_paired_both_review/raw.csv. The compact summary.csv is derived from those raw rows, while manifest.json records the workload count, methods, timing boundary, and host scope. The raw rows are the primary evidence for the reported medians; the summary is a convenience for reading them."),
            ("p", "Training, holdout, and paired evaluation use disjoint seed ranges: 11-16 for fitting, 101-104 for label holdout, 7001-7005 for the five general benchmark patterns, and 11001-11010 for late-heavy trials. A seed fixes task generation and the method-order shuffle; it does not freeze operating-system timing. Fresh elapsed seconds can move with host load, so the report keeps the saved paired experiment as its numeric reference. Every result is tied to a generated pattern rather than a claim about an unseen production trace."),
            ("h", "4.2 Why these workload patterns matter"),
            ("p", "Uniform tasks are a control for dispatch overhead: there is little imbalance to correct. Moderate and high patterns introduce different degrees of cost variation without changing the task count. Phase shift puts different work distributions in the two halves of a run, which tests whether online feedback is useful after conditions change. Prime scan substitutes a genuine prime-counting C kernel for the synthetic iteration loop, while the deliberately late-heavy pattern exposes the cost of leaving an unusually expensive task near the end. Together they vary one important scheduling property at a time, but they still cover only a narrow part of real multiprocessing work."),
        ],
    ),
    (
        "Experimental plan and measures",
        [
            ("h", "5.1 Software and file roles"),
            ("p", "The native/work_kernel.c file defines the two deterministic CPU functions. The task generators live in os_sched/workload.py and os_sched/straggler_workload.py. os_sched/scheduler_v2.py contains the fixed and online-tree paths, timing capture, and batch-result fields. os_sched/tree.py contains the small decision-tree implementation; os_sched/scheduler_v4.py contains cost estimation and balanced batching. The paired experiment and its independent audit are in os_sched/paired_both_review.py and os_sched/verify_paired.py. These files separate task computation, scheduling policy, and evaluation rather than treating a chart or screenshot as the implementation."),
            ("p", "Bash and GCC compile the C library, while Python 3.12.3 creates multiprocessing workers and handles CSV and JSON records. The code uses the Python standard library for the tree rather than scikit-learn, so model fitting and live inference have no additional Python package dependency. Linux under WSL2 is the reported host. The process-pool behavior and its overhead are part of the experiment; this is not a Linux kernel patch or a direct comparison of kernel schedulers."),
        ],
    ),
    (
        "Measured comparison and interpretation",
        [
            ("h", "6.1 Pairing and metric calculation"),
            ("p", "For each of the 35 workload-and-seed groups, the code constructs one task list and runs fixed 4, the loaded tree, and balanced 32 on that same list. Method order is deterministically shuffled within each group to avoid always giving one method the same position. The experiment rejects a group if the three final checksums differ. It produces 105 raw method rows and 18 workload-by-method summary rows. The audit independently checks row count, method coverage and order, checksum equality, and recomputed medians. No timed run is replaced by a made-up value when explaining the chart."),
            ("p", "The primary statistic for method M on one workload is median over seeds of time(M, seed) divided by time(fixed 4, same seed). Thus 0.870 means that the typical paired ratio is about 13 percent below the fixed-4 reference; 1.245 means about 24.5 percent above it. A ratio of displayed median times is not generally equal to the median of paired ratios. The latter retains seed matching and is the value used for the comparison in Table 1 and Figure 3."),
            ("h", "6.2 Why latency measures need care"),
            ("p", "Batch p95 is calculated from the response times of batches within a run, from submission to completion, and the report displays the median of those per-run p95 values across seeds. It describes a slow portion of batch responses, not the elapsed time of the whole job. Startup time describes the period before workers are ready; queue wait and worker service can also be inspected in raw.csv. These components are useful for diagnosis, but their medians or percentiles cannot simply be added to reconstruct the full wall-time median."),
            ("p", "The full wall boundary starts before planning and pool startup and ends after all results return and the pool closes. This prevents a method from appearing fast merely because its setup or teardown was timed elsewhere. Offline tree training is excluded from an invocation because the saved model is reused, but its cost and data collection are documented separately. All measurements remain exploratory: five seeds for each general pattern, ten for late-heavy, one WSL2 host, and two workers. The report does not attach confidence intervals or assert that these values will transfer to another machine."),
        ],
    ),
    (
        "Implementation progress and valid outputs",
        [
            ("h", "7.1 Reading the workload-level results"),
            ("p", "The tree's best paired ratio in the present study is on prime scan, 0.894, while its high-variance ratio is 1.245. That uneven outcome is important: the model may choose useful chunks under one cost pattern yet lose time through repeated round boundaries under another. Cost balancing's strongest intended result is late-heavy at 0.870. Its other paired ratios are below one on uniform, moderate, phase shift, and prime scan, but high variance is slightly above one at 1.006. These are workload-specific observations, not a claim that either path is always superior."),
            ("p", "The late-heavy diagnostic isolates a common interpretation error. The tree records a lower batch p95 than fixed 4 or cost balancing, yet its full wall time is higher than both in Table 2. The path repeatedly waits for complete rounds; batch response behavior alone does not account for idle time, dispatch sequencing, planning, and shutdown. Cost balancing has a higher batch p95 on this pattern but a lower full wall time because its estimated expensive batch starts early. This is a plausible mechanism supported by the implementation and live trace; the saved aggregate data do not isolate one causal component experimentally."),
            ("h", "7.2 What this contributes and what it does not"),
            ("p", "Adaptive self-scheduling, ML-assisted task scheduling, and work-stealing already appear in the cited literature. This prototype's contribution is a reproducible, same-task comparison of a small interpretable online tree, a single-dispatch cost-balancing policy, and a fixed reference while accounting for complete application wall time. The evidence shows the two alternatives have different bottlenecks. A potentially useful combination is a hypothesis for the next stage, not a worldwide-first novelty claim and not a measured Review 2 win."),
            ("h", "7.3 Why combine the two paths"),
            ("p", "The decision tree uses measured pilot behavior rather than relying on task metadata. It beats fixed 4 on prime scan (paired ratio 0.894), but loses on the other five patterns, most on high variance (1.245). Prediction is quick; the current scheduler waits for each round to finish before deciding again. Those barriers may idle workers, though their exact contribution to the loss is unmeasured."),
            ("p", "Cost balancing makes one dispatch plan and starts estimated expensive batches early. It wins most on the constructed late-heavy pattern (0.870) and is below 1 on four other patterns. It is slightly slower on high variance (1.006). It cannot correct bad cost estimates from live measurements; wrongly ranked tasks could erase its benefit. That failure mode still needs testing."),
            ("p", "The proposed hybrid would use one pilot and one tree prediction to choose a batch-size target, then cost-balance and order only unstarted tasks without repeated round waits. A gate would reject cost estimates that disagree with pilot service times; little remaining work would fall back to fixed dispatch. This aims to retain adaptive sizing and early heavy-task execution while limiting barrier and estimation costs. It may not succeed: no hybrid wall time has been measured."),
        ],
    ),
    (
        "Challenges and remaining work",
        [
            ("h", "8.1 What the present implementation demonstrates"),
            ("p", "The generator and C kernel execute repeatably; the trainer produces a serialized tree and a separate holdout score; the online scheduler loads that tree and logs chunk changes; the cost planner produces balanced batches and moves the estimated heavy task early; and the evaluator writes raw and summary metrics that can be audited. The terminal evidence shows live tests, a saved-data audit, and a fresh run of both scheduling paths. The screenshots establish that the modules ran in the reported environment, while raw.csv and the source code are needed to inspect the full multi-seed evidence."),
            ("p", "The progress count of six delivered modules out of seven is a functional inventory, not a measured percentage of engineering effort. The combined scheduler has no implemented path or measured row in the current paired results. The report therefore separates demonstrations of existing modules from ideas scheduled for later work. A reviewer can trace each table and figure back to the saved output files and compare the claimed method boundary with the source."),
        ],
    ),
    (
        "Terminal results",
        [
            ("h", "9.1 A testable next-stage design"),
            ("p", "The next implementation would run a short pilot, compute the existing five tree features, and make one tree prediction for the remaining work. It would translate the predicted chunk into a target batch cost, group unstarted tasks by estimated cost, and submit the heaviest batches first without another full-round wait. Before using that plan, it would compare pilot task-cost estimates with measured service times and check whether enough work remains to repay planning overhead. If either check fails, it would dispatch the remaining tasks with fixed chunk 4. The agreement and overhead thresholds must be set on development runs, not adjusted after seeing the final test results."),
            ("p", "A valid test would run the proposed combined path beside fixed 4, tree-only, and balance-only under the same generated task lists and shuffled method order. It should report checksum equality, full wall time and paired ratios first, then use startup, queue wait, service, and batch-tail measures to explain the result. Separate ablations should remove the gate or the balancing step to show whether either part contributes. More workers, larger task counts, perturbed cost estimates, a second Linux host, and a non-synthetic workload would test how far the conclusion travels. None of those future measurements is inserted into the present result tables."),
            ("h", "9.2 Practical limitations"),
            ("p", "The six generated patterns cannot represent arbitrary OS workloads, and the late-heavy pattern was chosen precisely because a metadata-based ordering may help it. Only two worker processes and one host were used; scaling and NUMA effects are unmeasured. The prime kernel is a useful distinct computation but still not a trace from an external application. The cost estimator can be wrong, and the tree's profiling labels are noisy and only 73.3 percent exact on holdout. These limitations make the most defensible conclusion narrow: the current implementation works and exposes a meaningful tradeoff, but general speedup and the value of a hybrid remain unproven."),
        ],
    ),
]

APPENDIX = [
    ("h", "Evidence files and reproduction notes"),
    ("p", "The screenshots above serve different purposes. Figure A1 records a fresh twelve-test Linux run; it checks functions and invariants, not scheduler speed. Figure A2 runs the saved-data audit, which counts 35 matched input groups and 105 timed method rows, verifies checksum equality, and recalculates the reported paired ratios. Figure A3 shows two live demonstrations, so its elapsed seconds are not substituted for the multi-seed medians in Tables 1 and 2. The figures make the outputs easy to show during review, while the CSV and JSON files provide the inspectable records behind the result claims."),
    ("p", "For the tree trace, the first decision begins at chunk 4 and selects chunk 16 after the 16-task pilot. The screenshot includes both the tree and guarded-tree first decisions, plus an END line that confirms equal checksums against fixed 4. The balanced trace reports that late task 187 was moved to dispatch position 0 and again confirms equal checksums. These are observations from the particular live seeds shown, not an effect size. The paired study instead evaluates the declared seed ranges and reports medians over five or ten matched trials."),
    ("p", "The repository keeps the executable source in native/ and os_sched/, automated checks in tests/, the trained model in results_final_linux/model.json, and the Review 2 records in results_paired_both_review/. In that last directory raw.csv contains workload, seed, method order, elapsed time, checksum, batch p95, startup, and other timing fields for each run. summary.csv reduces those rows to workload-by-method medians; manifest.json records the scope and timing boundary. The PDF and editable Word report are in output/. The repository link below leads to those artifacts and their source."),
    ("p", "A rebuild compiles the C library before Python tests or experiments. A fresh run of the paired experiment may produce different seconds because WSL2 and the host schedule processes under changing load; it should nevertheless preserve workload definitions, method set, checksum equality, row counts, and the audit procedure. Claims in this report refer to the saved verified run, not to an assurance that every rerun will keep the same ranking. This is why the exact task seeds, method ordering rule, raw measurements, and verifier are published together."),
]


def find_anchor(doc, phrase):
    for p in doc.paragraphs:
        if phrase in p.text and p.style is not None and p.style.name == "Heading 1":
            return p
    raise ValueError(f"Heading not found: {phrase}")


def insert_content(anchor, blocks):
    for kind, content in blocks:
        paragraph = anchor.insert_paragraph_before(content)
        if kind == "h":
            paragraph.style = "Heading 2"
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0, 0, 0)


def main():
    doc = Document(SOURCE)
    if "Heading 2" not in doc.styles:
        subheading = doc.styles.add_style("Heading 2", WD_STYLE_TYPE.PARAGRAPH)
        subheading.base_style = doc.styles["Normal"]
    original = [p.text for p in doc.paragraphs]
    for heading, blocks in INSERTIONS:
        insert_content(find_anchor(doc, heading), blocks)
    figure_two_caption = next(p for p in doc.paragraphs if p.text.startswith("Figure 2  The two implemented scheduling paths"))
    figure_two_caption.paragraph_format.keep_with_next = True
    references = next(p for p in doc.paragraphs if p.text.strip() == "References")
    insert_content(references, APPENDIX)
    iterator = iter(p.text for p in doc.paragraphs)
    assert all(any(text == source for text in iterator) for source in original), "Original text was changed"
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
