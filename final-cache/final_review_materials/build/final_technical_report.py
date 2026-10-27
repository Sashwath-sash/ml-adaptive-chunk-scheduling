"""Build the evidence-led, portrait Review 3 technical companion."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT = Path(__file__).resolve().parents[2]
MAT = PROJECT / "final_review_materials"
OUTPUT = MAT / "output"
PLOTS = OUTPUT / "review_plots"
DEST = OUTPUT / "Review3_Project_Technical_Report.docx"
diag = json.loads((OUTPUT / "frozen_model_diagnostics.json").read_text())
arc = json.loads((PROJECT / "results_hybrid_native_benchmark.json").read_text())
arc.update(json.loads((PROJECT / "results_calgary_hybrid_arc_5runs.json").read_text()))
s3 = json.loads((PROJECT / "results_hybrid_s3_native_benchmark.json").read_text())

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.left_margin = sec.right_margin = Inches(.78)
sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.62)
sec.header_distance = Inches(.3)
sec.footer_distance = Inches(.34)

for name, size, bold in [("Normal", 11, False), ("Title", 17, True),
                         ("Subtitle", 11, False), ("Heading 1", 14, True),
                         ("Heading 2", 12, True), ("Caption", 9, False)]:
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(5 if name == "Normal" else 7)
    if name == "Normal":
        st.paragraph_format.line_spacing = 1.12
    if name.startswith("Heading"):
        st.paragraph_format.keep_with_next = True
title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
title_borders = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_borders.append(title_bottom)
title_ppr.append(title_borders)

def para(text="", *, style=None, justify=True):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def heading(text, level=1):
    doc.add_heading(text, level=level)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    if widths:
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Inches(w)
    for ri, row in enumerate(t.rows):
        if ri == 0:
            trpr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), "90")
                node.set(qn("w:type"), "dxa")
                mar.append(node)
            tcpr.append(mar)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7E7E7")
                tcpr.append(shade)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{side}")
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), "4")
                edge.set(qn("w:color"), "B8B8B8")
                borders.append(edge)
            tcpr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
                    if ri == 0:
                        run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def figure(path, number, caption, explanation, *, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(style="Caption")
    c.paragraph_format.keep_with_next = True
    c.add_run(f"Figure {number}. {caption}")
    explanation_p = para(explanation)
    explanation_p.paragraph_format.keep_together = True

def new_page():
    doc.add_page_break()

# Start directly with the project title, with no separate cover page.
doc.add_paragraph("Machine Learning Assisted Adaptive Chunk Caching for IO Workloads", style="Title")
para("Review 3 technical report and experimental evidence | Sashwath J | 24BAI0059", style="Subtitle", justify=False)
heading("Abstract")
para("This project tests whether a learned cache-admission decision can reduce backing-file reads in a bounded chunk cache. A histogram gradient-boosting classifier predicts whether a requested key will reappear within the next 1,024 trace requests. On a cache miss, a native C replay reads the 4 KiB block, evaluates the frozen model, and either bypasses insertion or passes the block to ARC or a local S3-FIFO replacement policy. At a 512-block capacity, the ML-assisted variants made fewer aligned direct-read calls than their respective parent policies in all eight saved comparisons across four held-out streams. The reduction ranged from about 1.0% to 7.5%. Five-run median elapsed time improved in six comparisons and worsened in two. The supported outcome is therefore improved read-call efficiency in this prototype, with a measurable compute and timing trade-off. The result does not establish physical-device savings or a universal latency gain.")
para("Keywords: cache admission; chunk caching; gradient boosting; direct I/O; ARC; S3-FIFO; trace replay.")

heading("1 Introduction and scope transition")
para("The Review 2 implementation investigated CPU chunk scheduling for multiprocessing workloads. Its tests did not establish a reliable overall completion-time benefit. The final prototype keeps the larger question of learned decisions about fixed-size work units, but moves the decision to I/O cache admission: whether a fetched chunk should occupy limited cache capacity. This is a disclosed change of application and metric, not an improvement to the old scheduler. The CPU results are not combined with the I/O results below.")
para("A cache miss requires the program to fetch a backing block. In a small cache, admitting a block unlikely to be reused may evict one that would have served a later request. The model supplies a future-reuse estimate from information available at the current request. The native cache still serves every miss correctly; the prediction changes only retention. The primary metric is the number of aligned direct-read calls made while serving a fixed trace. Elapsed time is a secondary check because model evaluation also costs CPU time.")

heading("2 Problem statement objectives and contribution")
para("Problem statement. Given a bounded 512-entry cache of 4 KiB chunks and a request stream, can a frozen prediction of near-term reuse reduce backing-file read calls when placed before established replacement policies?")
para("The implementation objectives are to train a model on chronological trace data without future information in online features; export it to an actual C decision path; compare each ML-assisted cache with its unchanged parent policy on held-out streams; and verify equal outputs by checksum. The project-level contribution is this measured integration and its documented transfer and timing limits. ML admission and the replacement policies themselves already have prior art, so the report does not claim invention of a new general algorithm.")

heading("3 Prior work context")
para("CACHEUS uses learned adaptation among cache experts, which establishes that machine learning for cache replacement is not new (Rodriguez et al., 2021). S3-FIFO studies a simple three-queue replacement design, including an admission-like small queue that limits pollution by one-use objects (Yang et al., 2023). Baleen combines ML admission and prefetching for flash caches and evaluates an end-to-end backend-load metric (Wong et al., 2024). Our prototype asks a narrower engineering question: what happens when a frozen reuse classifier filters admissions before two existing parent policies in one native direct-I/O replay? These three anchors are context, not the 30-paper analytical survey required for the student's final paper.")

heading("4 System architecture and module workflow")
figure(OUTPUT / "cache_data_flow.png", 1,
       "Chronological training, validation, model selection and held-out native replay.",
       "BrowseTrace A and B and the earlier NASA July web log supply chronological training and validation requests. Candidate models are fitted from the earlier requests; validation miss counts select the saved configuration. Separate NASA August and Calgary web logs join the final test inputs, not the training inputs. The frozen tree nodes are exported to C before the held-out replay. This branch and feedback loop is the actual model-selection procedure, not a runtime learning loop.", width=6.5)
figure(OUTPUT / "cache_request_flow.png", 2,
       "Cache hit, cache miss, model admission and next-request loop.",
       "A hit returns the resident 4 KiB block without a backing read. A miss requires an aligned read; the frozen model then scores eight current-or-past features. A score below 0.2 bypasses insertion, while a score at or above 0.2 passes the block to ARC or local S3-FIFO, which may evict a resident block. Both paths return the same data and update the checksum before processing the next request. The model changes future retention, not the read needed for the current miss.", width=6.5)

new_page()
heading("5 Data and training procedure")
para("All sources in this report are sequences of web requests, not spacecraft measurements or operating-system sensor data. BrowseTrace A and B contribute chronological 60% training, 20% validation and final 20% test slices. Earlier requests from a 1995 NASA website access log add training and validation examples. A later August log from the same NASA website tests transfer over time. The Calgary website access log tests transfer to a different server. Only the last 100,000 valid successful GET requests from each separate test log are replayed. The training set contains 414,142 requests and the validation set contains 138,047. The reuse label is one if the same key occurs again within the next 1,024 requests in that trace slice; the label is not available to the live cache.")
table(["Source", "Train", "Validation", "Final test"], [
    ["BrowseTrace A", "49,473", "16,491", "16,491"],
    ["BrowseTrace B", "214,669", "71,556", "71,557"],
    ["NASA web July", "150,000", "50,000", "Excluded tail"],
    ["NASA web August", "0", "0", "100,000"],
    ["Calgary web", "0", "0", "100,000"],
], [2.05, 1.18, 1.18, 1.62])
para("Table 1. Request counts by source and use. The NASA July tail was excluded after an earlier transfer failure. The ML+S3 pairing was explored after the Calgary plain-policy result was viewed, so Calgary is not a blind selection test of that pairing.")

heading("6 Frozen model and online features")
para("The actual fitted estimator is scikit-learn's HistGradientBoostingClassifier. It has 100 boosting stages, at most 15 leaves per tree, at least 40 samples per leaf, learning rate 0.07, L2 regularization 2 and random seed 811. The exported model contains 2,900 numerical nodes. A validation search over 40 or 100 stages and thresholds 0.2, 0.4 and 0.6 selected the saved 100-stage, 0.2-threshold configuration. No retraining takes place during native replay.")
table(["Online feature", "Measured from", "Purpose"], [
    ["Log response size", "Current object metadata", "Object-size signal"],
    ["Log prior key count", "Earlier same-key requests", "Frequency"],
    ["Log recency", "Distance to prior same-key request", "Recent locality"],
    ["Log prior host count", "Earlier host requests", "Host activity"],
    ["Path length", "Current URL path", "Resource structure"],
    ["Static suffix", "Current path suffix", "Resource-type proxy"],
    ["Query flag", "Current URL", "Dynamic URL proxy"],
    ["Key length", "Current full key", "Key structure"],
], [1.65, 2.35, 2.05])
para("Table 2. Eight online inputs. The live program does not inspect the future reuse label. The special recency value for a never-seen key uses the known replay-slice length; this is a replay convention that should become a fixed sentinel in a production service.")

heading("7 Native implementation and correctness checks")
para("The C loop keeps at most 512 resident 4 KiB blocks, equal to 2 MiB of data capacity. A miss issues an aligned O_DIRECT read to a deterministic surrogate backing file, evaluates the exported trees and optionally inserts the block. The baseline ARC and local S3-FIFO paths use the same request order, cache capacity, block size and backing data. Each variant returns the same checksum for a given stream. Independent Python reference replays match the native miss counts, and saved evidence includes sanitizer checks. The backing blocks are surrogates because the original web response bodies are unavailable; a read call is therefore the precise measured I/O event, not a claim about production HTTP latency or physical disk transactions.")

new_page()
heading("8 Evaluation metrics and comparison design")
para("Direct-read calls count how often the native program requests a backing block; fewer calls mean better cache efficiency in this replay. Hit rate is hits divided by all requests. The per-run loop time is measured inside the native request loop. Five runs were stored for each policy and stream, and the report uses their median as a descriptive summary. The range is shown because five single-host observations are insufficient for a stable performance guarantee. A fair ML effect compares ML+ARC against ARC and ML+S3-FIFO against S3-FIFO on the same stream. Comparing the two parent algorithms alone does not isolate the learned admission decision.")
table(["Stream", "ARC", "ML + ARC", "S3-FIFO", "ML + S3-FIFO"], [
    ["BrowseTrace A", "7,430", "7,104", "7,773", "7,336"],
    ["BrowseTrace B", "17,913", "17,740", "18,950", "18,586"],
    ["NASA web", "8,792", "8,130", "8,349", "8,048"],
    ["Calgary web", "16,493", "15,718", "15,406", "15,177"],
], [1.65, 1.0, 1.2, 1.15, 1.35])
para("Table 3. Exact saved native direct-read calls. Both ML variants made fewer calls than their respective parents on each stream. The next section shows the same results visually and keeps timing separate.")

heading("9 Experimental figures and interpretation")
para("Each figure below is generated from saved benchmark JSON or frozen-model diagnostics. The explanation immediately following it states both what the plot measures and what it cannot establish. The plotting source and tabulated values remain in the project for inspection.")

gallery = [
    ("01_direct_read_calls.png", "Aligned direct-read calls across four held-out streams.",
     "These are absolute 4 KiB read-call counts, not inferred disk transactions. On BrowseTrace A, for example, ML+ARC makes 7,104 calls versus ARC's 7,430, while ML+S3-FIFO makes 7,336 versus its parent's 7,773. The same within-parent direction holds on all four streams. The different heights of ARC and S3-FIFO alone should not be attributed to the model."),
    ("02_read_reduction_percent.png", "Percentage read reduction against each parent policy.",
     "The percentage calculation is (parent reads - hybrid reads) / parent reads. The largest saved reduction is about 7.5% for ML+ARC on the NASA August web log; the smallest is about 1.0% for ML+ARC on BrowseTrace B. Every bar is positive, but the size of the effect changes with the request stream and replacement policy. This supports a modest, workload-dependent I/O-efficiency benefit."),
    ("03_timing_all_runs.png", "All five saved timing observations for each policy and stream.",
     "Dots are individual native loop times and the short vertical mark is their median. The wide spread in several groups, especially Calgary-web ARC and NASA-web S3-FIFO, shows why one fast run would be misleading. These points are not confidence intervals and were not collected on multiple hosts. The precise direct-read counts repeat, while wall time is sensitive to execution conditions."),
    ("04_median_time_change.png", "Change in five-run median loop time after adding ML admission.",
     "Six of eight bars lie above zero, meaning the stored median loop time decreased. Two lie below: ML+ARC on BrowseTrace B and ML+S3-FIFO on the NASA August web log. The latter is about 4% slower despite fewer reads. The experiment does not isolate model-evaluation CPU time from all other timing variation, so the chart supports a trade-off observation, not a proven causal breakdown."),
    ("05_reads_vs_time_tradeoff.png", "Read-call savings and elapsed-time change in the same comparison.",
     "Points to the right save backing reads. Points above zero also have lower median loop time. Two points are right of zero but below zero, showing that fewer read calls need not make a run faster. The case for this system is reducing expensive read requests when that objective matters; it cannot be sold as an unconditional speedup."),
    ("06_cache_hit_rates.png", "Measured cache-hit rates for four policies on four streams.",
     "Hit rate equals recorded hits divided by hits plus direct-read misses. The ML-filtered version has a slightly higher hit rate than its parent in each saved pairing, consistent with the read counts. The figure also shows that baseline difficulty differs by stream: the NASA August web log has much higher hit rates than BrowseTrace A, so identical percentage gains should not be expected."),
    ("07_model_quality_and_prevalence.png", "Held-out ranking quality and target prevalence for the frozen model.",
     "ROC AUC ranges from about 0.816 on BrowseTrace B to 0.962 on the NASA August web log. Average precision must be read alongside prevalence because reuse is common in these slices. A high classifier score describes discrimination of the defined 1,024-request reuse label; it does not by itself prove a cache-level read saving. That claim comes from the native replay."),
    ("08_precision_recall_at_threshold.png", "Precision and recall at the operational 0.2 threshold.",
     "The permissive threshold retains nearly all positive reuse labels, with recall about 0.982 to 0.997. Precision is lower, ranging about 0.720 to 0.885. In plain terms, the model rarely rejects a chunk that will reappear in the defined horizon, but it still admits some chunks that do not. These values use offline future labels only for evaluation."),
    ("09_classification_outcomes.png", "Classification outcomes as a proportion of each held-out stream.",
     "The false-positive portion represents admitted blocks that do not reappear within the labeling horizon. The false-negative portion is small because the threshold favors recall. Class prevalence differs by stream, so raw TP and FP proportions should not be compared without the positive-share context in Figure 9. The finite trace end right-censors some labels."),
    ("10_early_stream_misses.png", "Cumulative misses at the first three recorded checkpoints.",
     "These exact checkpoints occur after 1,024, 2,048 and 4,096 requests. The line segments connect only those observations; they are not a continuous trace of every intervening miss. Early differences are sometimes small even when the full-stream totals diverge. This figure therefore documents warm-up behavior but cannot replace the complete-count comparison in Figure 3."),
]
for i, (file, caption, explanation) in enumerate(gallery, start=3):
    heading(f"9.{i-2} {caption[:-1]}", 2)
    figure(PLOTS / file, i, caption, explanation, width=6.2)

heading("10 Model calibration and interpretation")
figure(OUTPUT / "frozen_model_diagnostics_bw.png", 13,
       "ROC and calibration curves recomputed from the saved model on final test slices.",
       "The ROC panel tests whether higher scores rank future-reuse examples above non-reuse examples. The calibration panel compares predicted probability with observed frequency. The curves differ by source, which is evidence of distribution shift and imperfect probability calibration. The admission threshold was chosen on validation data, not retuned on these held-out slices. The Calgary web log is not fully blind for the later ML+S3 pairing, as stated earlier.", width=6.2)
table(["Test stream", "ROC AUC", "Precision", "Recall", "Brier"], [
    ["BrowseTrace A", ".862", ".720", ".993", ".143"],
    ["BrowseTrace B", ".816", ".744", ".997", ".150"],
    ["NASA web", ".962", ".885", ".991", ".058"],
    ["Calgary web", ".889", ".805", ".982", ".118"],
], [1.65, 1.05, 1.05, 1.05, 1.05])
para("Table 4. Frozen-model diagnostics at threshold 0.2 except ROC AUC and Brier, which use scores over all thresholds. These are classifier measures, not direct system metrics.")

heading("11 Interpretation limitations and remaining research")
para("The evidence supports fewer native direct-read calls with learned admission on the four tested request streams, under a fixed 512-block cache and deterministic 4 KiB backing data. It does not show that physical device I/O, energy, wear, network latency or multiprocess throughput improved. The measured request loop is single-process and user-space; its S3-FIFO branch is local rather than the official optimized implementation. The cache traces have finite-horizon labels and source shifts. Calgary was inspected before the later S3 pairing. Five timing repeats on one machine should be treated as indicative only. Larger caches, different chunk sizes and genuinely concurrent shared-cache workloads would be follow-up experiments, not claims of completed work.")

heading("12 Demonstration and reproducibility")
para("From the project directory, `python demo_hybrid_final.py` checks saved source hashes, model settings, read counts and policy checksums, then prints the four-stream comparison. `python demo_hybrid_final.py --live calgary` recompiles and runs a fresh native Linux direct-I/O replay through WSL and GCC. Fresh counts and checksums should match the saved result; elapsed times may differ. `python verify_hybrid_native.py` independently checks the native policies against Python reference replays. The plotting command is `.venv\\Scripts\\python.exe plot_review_gallery.py` on the Windows project host.")
para("The inspectable code paths are `train_three_source_cache.py` for training, `export_cache_native.py` for frozen model export, `cache_native.c` for online reads and admission, `benchmark_hybrid_native.py` for stored repeats, `evaluate_frozen_model.py` for classifier diagnostics, and `plot_review_gallery.py` for the numbered result figures. The JSON inputs and the exported model are retained alongside them. Load the trusted local `.joblib` model only; pickle-family model formats can execute code during loading.")

heading("13 Conclusion")
para("The prototype uses a real trained model in the native miss path and produces a small but consistent reduction in backing-file read calls against each of its two parent cache policies on four saved streams. The gain is not free: the program must evaluate the model, and two stored comparisons have worse median elapsed time. The defensible final result is improved cache-read efficiency for these traces and this capacity, with the implementation and failure cases available for inspection. It is a scope change from Review 2 CPU chunk scheduling, not a claim that the original scheduling problem was solved.")

heading("References")
refs = [
    "Rodriguez, L. V., Yusuf, F., Lyons, S., Paz, E., Rangaswami, R., Liu, J., Zhao, M., & Narasimhan, G. (2021). Learning cache replacement with CACHEUS. In 19th USENIX Conference on File and Storage Technologies (pp. 341-354). USENIX Association. https://www.usenix.org/conference/fast21/presentation/rodriguez",
    "Wong, D. L.-K., Wu, H., Molder, C., Gunasekar, S., Lu, J., Khandkar, S., Sharma, A., Berger, D. S., Beckmann, N., & Ganger, G. R. (2024). Baleen: ML admission and prefetching for flash caches. In 22nd USENIX Conference on File and Storage Technologies (pp. 347-371). USENIX Association. https://www.usenix.org/conference/fast24/presentation/wong",
    "Yang, J., Zhang, Y., Qiu, Z., Yue, Y., & Rashmi, K. V. (2023). FIFO queues are all you need for cache eviction. In Proceedings of the 29th Symposium on Operating Systems Principles. Association for Computing Machinery. https://github.com/Thesys-lab/sosp23-s3fifo",
]
for ref in refs:
    p = para(ref)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(.3)
    p.paragraph_format.first_line_indent = Inches(-.3)

heading("Academic submission status")
para("This is an illustrated technical evidence companion for review and student editing. The supplied Review 3 guideline requires an original student-written paper with at least 30 relevant peer-reviewed sources from 2021-2026, analytical comparison and fully matched APA citations. This document has three verified context sources and does not meet that paper requirement by itself. The student must complete and verify the literature work, write the submission in their own words, and confirm the changed scope with the reviewer before claiming final academic compliance.")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Review 3 technical evidence companion  |  Sashwath J  |  24BAI0059")
OUTPUT.mkdir(parents=True, exist_ok=True)
doc.save(DEST)
print(DEST)
